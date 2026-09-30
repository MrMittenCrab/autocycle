#!/usr/bin/env python3
"""Optional macOS integration probe; only creates/opens disposable workbooks."""
from pathlib import Path
import sys
import tempfile
import openpyxl
from excel_verification import verify
from adjudication import location

folder=Path(tempfile.mkdtemp(prefix='autocycle-excel-check-'))
source=folder/'disposable.xlsx'
wb=openpyxl.Workbook();ws=wb.active;ws.title='Summary'
ws['A1']='=Inputs!A1*2';ws['A2']='=A1+1'
wb.create_sheet('Inputs')['A1']='=20+1';wb.save(source);wb.close()
check=folder/'check.py'
check.write_text('''import openpyxl,sys
wb=openpyxl.load_workbook(sys.argv[1],data_only=True)
assert wb['Summary']['A1'].value==42, wb['Summary']['A1'].value
assert wb['Summary']['A2'].value==43, wb['Summary']['A2'].value
assert wb['Inputs']['A1'].value==21, wb['Inputs']['A1'].value
wb.close()
wb=openpyxl.load_workbook(sys.argv[1],data_only=False)
assert wb['Summary']['A1'].value=='=Inputs!A1*2'
wb.close()
print('Native cached values and preserved formula verified')
''')
result=verify(source,[sys.executable,str(check),'{workbook}'],location())
print(result)
sys.exit(0 if result['status']=='VERIFIED' else 2)
